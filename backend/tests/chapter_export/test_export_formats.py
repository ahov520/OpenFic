# -*- coding: utf-8 -*-
"""导出格式参数化测试：EPUB/DOCX 产物读回、元数据与清理行为。"""

from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub

from app.chapter_export import service as chapter_export_service
from app.chapter_export import formats as export_formats
from app.chapter_export.docx_writer import write_docx
from app.chapter_export.epub_writer import write_epub
from app.chapter_export.formats import (
    ExportDocumentChapter,
    ExportDocumentSpec,
    ExportDocumentVolume,
)
from app.background.jobs.models import BackgroundJob
from app.chapter_export.service import cleanup_chapter_export_files, export_file_paths


def _sample_spec() -> ExportDocumentSpec:
    return ExportDocumentSpec(
        book_title="测试小说",
        author="OpenFic",
        language="zh",
        volumes=[
            ExportDocumentVolume(
                title="第一卷 风起",
                order=1,
                chapters=[
                    ExportDocumentChapter(title="第一章", content="第一章正文\n第二行"),
                    ExportDocumentChapter(title="第二章", content="第二章正文"),
                ],
            ),
            ExportDocumentVolume(
                title="第二卷 云涌",
                order=2,
                chapters=[ExportDocumentChapter(title="第三章", content="第三章正文")],
            ),
        ],
    )


def test_epub_writer_roundtrip(tmp_path) -> None:
    path = tmp_path / "book.epub"
    write_epub(path, _sample_spec())

    book = epub.read_epub(str(path))
    assert book.get_metadata("DC", "title")[0][0] == "测试小说"
    assert book.get_metadata("DC", "creator")[0][0] == "OpenFic"
    assert book.get_metadata("DC", "language")[0][0] == "zh"

    chapter_documents = [
        item
        for item in book.get_items_of_type(ITEM_DOCUMENT)
        if item.file_name.startswith("chapter_")
    ]
    assert len(chapter_documents) == 3
    volume_documents = [
        item
        for item in book.get_items_of_type(ITEM_DOCUMENT)
        if item.file_name.startswith("volume_")
    ]
    assert len(volume_documents) == 2

    first_chapter_content = chapter_documents[0].get_content().decode("utf-8")
    assert "第一章" in first_chapter_content
    assert "第一章正文" in first_chapter_content
    assert "第二行" in first_chapter_content


def test_docx_writer_roundtrip(tmp_path) -> None:
    path = tmp_path / "book.docx"
    write_docx(path, _sample_spec())

    document = Document(str(path))
    assert document.core_properties.title == "测试小说"
    assert document.core_properties.author == "OpenFic"
    assert document.core_properties.language == "zh"

    level1 = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading 1")
    ]
    level2 = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading 2")
    ]
    assert level1 == ["第一卷 风起", "第二卷 云涌"]
    assert level2 == ["第一章", "第二章", "第三章"]
    body_texts = [paragraph.text for paragraph in document.paragraphs]
    assert "第一章正文" in body_texts
    assert "第二行" in body_texts
    assert "第二章正文" in body_texts


def test_export_file_paths_follow_format(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    part_txt, output_txt = chapter_export_service.export_file_paths("job-1", "txt")
    part_epub, output_epub = chapter_export_service.export_file_paths("job-1", "epub")
    part_docx, output_docx = chapter_export_service.export_file_paths("job-1", "docx")

    assert output_txt.suffix == ".txt"
    assert output_epub.suffix == ".epub"
    assert output_docx.suffix == ".docx"
    assert part_epub == part_docx == part_txt
    assert part_txt.suffix == ".part"


def test_export_media_types_cover_all_formats() -> None:
    for export_format in export_formats.EXPORT_FORMATS:
        assert export_format in export_formats.EXPORT_MEDIA_TYPES
    assert export_formats.EXPORT_MEDIA_TYPES["epub"] == "application/epub+zip"
    assert "wordprocessingml" in export_formats.EXPORT_MEDIA_TYPES["docx"]


def _make_job(job_id: str, status: str, payload_format: str) -> BackgroundJob:
    return BackgroundJob(
        id=job_id,
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status=status,
        payload_json='{"format":"%s","filename":"书.%s"}' % (payload_format, payload_format),
        result_json='{"expires_at":"%s"}'
        % (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
    )


@pytest.mark.asyncio
async def test_cleanup_keeps_running_job_part_and_product_for_all_formats(
    session, monkeypatch, tmp_path
) -> None:
    """进行中的任务：.part 与成品（任意格式）都保留。"""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    for export_format in ("epub", "docx"):
        job = _make_job(f"running-{export_format}", "running", export_format)
        session.add(job)
        await session.commit()
        part_path, output_path = export_file_paths(job.id, export_format)
        part_path.write_bytes(b"part")
        output_path.write_bytes(b"product")

    removed = await cleanup_chapter_export_files(session)

    assert removed == 0
    for export_format in ("epub", "docx"):
        part_path, output_path = export_file_paths(f"running-{export_format}", export_format)
        assert part_path.exists()
        assert output_path.exists()


@pytest.mark.asyncio
async def test_cleanup_deletes_cancelled_job_files_and_expired_products(
    session, monkeypatch, tmp_path
) -> None:
    """取消任务的文件全部删除；成功任务过期成品删除。"""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)

    cancelled = _make_job("cancelled-epub", "cancelled", "epub")
    cancelled.result_json = "{}"
    session.add(cancelled)
    expired = _make_job("expired-docx", "succeeded", "docx")
    expired.result_json = '{"expires_at":"2020-01-01T00:00:00+00:00"}'
    session.add(expired)
    await session.commit()

    for job_id, export_format in (("cancelled-epub", "epub"), ("expired-docx", "docx")):
        part_path, output_path = export_file_paths(job_id, export_format)
        part_path.write_bytes(b"part")
        output_path.write_bytes(b"product")

    removed = await cleanup_chapter_export_files(session)

    assert removed == 4
    for job_id, export_format in (("cancelled-epub", "epub"), ("expired-docx", "docx")):
        part_path, output_path = export_file_paths(job_id, export_format)
        assert not part_path.exists()
        assert not output_path.exists()


@pytest.mark.asyncio
async def test_cleanup_keeps_fresh_epub_product_of_succeeded_job(
    session, monkeypatch, tmp_path
) -> None:
    """成功任务的 EPUB 成品在有效期内保留。"""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job = _make_job("fresh-epub", "succeeded", "epub")
    session.add(job)
    await session.commit()
    _part_path, output_path = export_file_paths("fresh-epub", "epub")
    output_path.write_bytes(b"epub-bytes")

    removed = await cleanup_chapter_export_files(session)

    assert removed == 0
    assert output_path.exists()
