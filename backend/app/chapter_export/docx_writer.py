"""DOCX 导出 writer：按卷/章结构生成 Word 文档，中文元数据（书名/作者/语言 zh）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from app.chapter_export.formats import ExportDocumentSpec


def write_docx(path: str | Path, spec: ExportDocumentSpec) -> None:
    """按 spec 生成 DOCX 文件并写盘。

    卷标题映射为一级标题，章节标题映射为二级标题，正文逐段成段；
    文档属性写入书名、作者与语言（zh）。
    """
    document = Document()
    core_properties = document.core_properties
    core_properties.title = spec.book_title
    core_properties.author = spec.author
    core_properties.language = spec.language

    for volume in spec.volumes:
        if volume.title is not None:
            document.add_heading(volume.title, level=1)
        for chapter in volume.chapters:
            document.add_heading(chapter.title, level=2)
            for paragraph in chapter.content.split("\n"):
                if paragraph.strip():
                    document.add_paragraph(paragraph)

    document.save(str(path))
